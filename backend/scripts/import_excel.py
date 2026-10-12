"""
Script para importar datos desde un archivo Excel de seguimiento ISO 27001.
Uso: python scripts/import_excel.py [ruta_al_archivo.xlsx]
"""
import sys
import os
from pathlib import Path

# Agregar el directorio padre al path para importar app
sys.path.insert(0, str(Path(__file__).parent.parent))

import openpyxl
from sqlalchemy.orm import Session

from app.db import SessionLocal, engine
from app.models.policy import Policy
from app.models.control import Control
from app.models.requirement import Requirement


def import_policies(sheet, db: Session) -> dict:
    """Importa políticas desde la hoja 'Políticas'"""
    policies = {}
    
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row[0]:  # Saltar filas vacías
            continue
            
        code = str(row[0]).strip()
        title = str(row[1]).strip() if row[1] else ""
        description = str(row[2]).strip() if len(row) > 2 and row[2] else ""
        
        # Buscar o crear política
        policy = db.query(Policy).filter(Policy.code == code).first()
        if not policy:
            policy = Policy(code=code, title=title, description=description)
            db.add(policy)
            print(f"  ✅ Política creada: {code} - {title}")
        else:
            policy.title = title
            policy.description = description
            print(f"  🔄 Política actualizada: {code}")
        
        policies[code] = policy
    
    db.commit()
    return policies


def import_controls(sheet, db: Session) -> dict:
    """Importa controles ISO 27001 desde la hoja 'Controles'"""
    controls = {}
    
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
            
        code = str(row[0]).strip()
        title = str(row[1]).strip() if row[1] else ""
        description = str(row[2]).strip() if len(row) > 2 and row[2] else ""
        category = str(row[3]).strip() if len(row) > 3 and row[3] else ""
        
        control = db.query(Control).filter(Control.code == code).first()
        if not control:
            control = Control(code=code, title=title, description=description, category=category)
            db.add(control)
            print(f"  ✅ Control creado: {code} - {title}")
        else:
            control.title = title
            control.description = description
            control.category = category
            print(f"  🔄 Control actualizado: {code}")
        
        controls[code] = control
    
    db.commit()
    return controls


def import_requirements(sheet, db: Session, policies: dict, controls: dict):
    """Importa requisitos desde la hoja 'Requisitos'"""
    count = 0
    
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
            
        code = str(row[0]).strip()
        title = str(row[1]).strip() if row[1] else ""
        description = str(row[2]).strip() if len(row) > 2 and row[2] else ""
        frequency = str(row[3]).strip() if len(row) > 3 and row[3] else "Anual"
        policy_code = str(row[4]).strip() if len(row) > 4 and row[4] else ""
        control_code = str(row[5]).strip() if len(row) > 5 and row[5] else ""
        
        # Validar que la política exista
        if policy_code not in policies:
            print(f"  ⚠️  Requisito {code}: política {policy_code} no encontrada, omitiendo")
            continue
        
        # Buscar control (opcional)
        control = controls.get(control_code)
        
        requirement = db.query(Requirement).filter(Requirement.code == code).first()
        if not requirement:
            requirement = Requirement(
                code=code,
                title=title,
                description=description,
                frequency=frequency,
                policy_code=policy_code,
                control_id=control.id if control else None,
            )
            db.add(requirement)
            print(f"  ✅ Requisito creado: {code} - {title}")
        else:
            requirement.title = title
            requirement.description = description
            requirement.frequency = frequency
            requirement.policy_code = policy_code
            requirement.control_id = control.id if control else None
            print(f"  🔄 Requisito actualizado: {code}")
        
        count += 1
    
    db.commit()
    return count


def main():
    """Función principal de importación"""
    # Ruta del archivo Excel
    excel_path = sys.argv[1] if len(sys.argv) > 1 else "Seguimiento_Evidencias_ISO27001.xlsx"
    
    if not os.path.exists(excel_path):
        print(f"❌ Error: El archivo {excel_path} no existe")
        print(f"Uso: python scripts/import_excel.py [ruta_al_archivo.xlsx]")
        sys.exit(1)
    
    print(f"📂 Leyendo archivo: {excel_path}")
    wb = openpyxl.load_workbook(excel_path, read_only=True)
    
    # Verificar que existan las hojas necesarias
    required_sheets = ["Políticas", "Controles", "Requisitos"]
    for sheet_name in required_sheets:
        if sheet_name not in wb.sheetnames:
            print(f"❌ Error: La hoja '{sheet_name}' no existe en el archivo")
            sys.exit(1)
    
    db = SessionLocal()
    
    try:
        print("\n📋 Importando políticas...")
        policies = import_policies(wb["Políticas"], db)
        print(f"   Total: {len(policies)} políticas")
        
        print("\n📋 Importando controles...")
        controls = import_controls(wb["Controles"], db)
        print(f"   Total: {len(controls)} controles")
        
        print("\n📋 Importando requisitos...")
        req_count = import_requirements(wb["Requisitos"], db, policies, controls)
        print(f"   Total: {req_count} requisitos")
        
        print(f"\n✅ Importación completada exitosamente")
        print(f"   Políticas: {len(policies)}")
        print(f"   Controles: {len(controls)}")
        print(f"   Requisitos: {req_count}")
        
    except Exception as e:
        print(f"\n❌ Error durante la importación: {e}")
        db.rollback()
        sys.exit(1)
    finally:
        db.close()
        wb.close()


if __name__ == "__main__":
    main()
